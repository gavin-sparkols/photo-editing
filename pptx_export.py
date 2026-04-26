# -*- coding: utf-8 -*-
"""由 batch JSON + uploads 内底图生成 PPTX（每页一幻灯片，底图 + 字框文本）。"""
from __future__ import annotations

import io
import os
from typing import Any

EMU_PER_INCH = 914400


def px_to_emu(px: float) -> int:
    """自然像素 → EMU（按 96dpi，与前端 canvas 习惯一致）。"""
    return int(px / 96.0 * EMU_PER_INCH)


def _rgb(region: dict[str, Any]) -> tuple[int, int, int]:
    st = region.get("style") or {}
    rgb = st.get("text_color_rgb")
    if isinstance(rgb, (list, tuple)) and len(rgb) >= 3:
        return int(rgb[0]), int(rgb[1]), int(rgb[2])
    return 26, 32, 44


def _font_pt(region: dict[str, Any]) -> float:
    v = region.get("export_font_nat")
    if v is None:
        v = region.get("font_size_px")
    try:
        n = float(v)
    except (TypeError, ValueError):
        n = 12.0
    return max(6.0, min(120.0, n * 72.0 / 96.0))


def _read_image_size(path: str) -> tuple[int, int]:
    try:
        from PIL import Image as PILImage

        with PILImage.open(path) as im:
            return int(im.size[0]), int(im.size[1])
    except Exception:
        return 960, 540


def _clear_slide_shapes(slide: Any) -> None:
    for shape in list(slide.shapes):
        el = shape._element
        el.getparent().remove(el)


def build_pptx_bytes(
    *,
    pages: list[dict[str, Any]],
    background_files: list[str | None],
    upload_dir: str,
) -> bytes:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Pt
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("请安装 python-pptx: pip install python-pptx") from e

    if not pages:
        raise ValueError("pages 为空")

    prs = Presentation()

    def pick_blank_layout():
        for layout in prs.slide_layouts:
            if (layout.name or "").lower() == "blank":
                return layout
        return prs.slide_layouts[-1]

    blank_layout = pick_blank_layout()

    first_bg = background_files[0] if background_files else None
    if not first_bg:
        raise ValueError("缺少背景图文件名")
    p0 = os.path.join(upload_dir, os.path.basename(first_bg))
    if not os.path.isfile(p0):
        raise ValueError(f"找不到背景图: {first_bg}")
    w0, h0 = _read_image_size(p0)
    if w0 < 1 or h0 < 1:
        w0, h0 = 960, 540
    slide_w = px_to_emu(w0)
    slide_h = px_to_emu(h0)
    prs.slide_width = slide_w
    prs.slide_height = slide_h

    for idx, layout in enumerate(pages):
        bg_name = (
            background_files[idx]
            if idx < len(background_files)
            else None
        )
        if not bg_name:
            raise ValueError(f"第 {idx + 1} 页缺少背景图文件名")
        bg_path = os.path.join(upload_dir, os.path.basename(bg_name))
        if not os.path.isfile(bg_path):
            raise ValueError(f"找不到背景图: {bg_name}")

        w_px, h_px = _read_image_size(bg_path)
        if w_px < 1 or h_px < 1:
            w_px, h_px = w0, h0

        if idx == 0:
            slide = prs.slides[0]
            _clear_slide_shapes(slide)
        else:
            slide = prs.slides.add_slide(blank_layout)

        slide.shapes.add_picture(bg_path, 0, 0, width=slide_w, height=slide_h)

        regions = layout.get("regions") or []
        for reg in regions:
            bb = reg.get("bbox") or {}
            try:
                x = float(bb.get("x", 0))
                y = float(bb.get("y", 0))
                bw = float(bb.get("width", 1))
                bh = float(bb.get("height", 1))
            except (TypeError, ValueError):
                continue
            text = (reg.get("text") or "").replace("\n", " ").strip()
            if not text:
                continue

            left = int(x / w_px * slide_w)
            top = int(y / h_px * slide_h)
            width = max(1, int(bw / w_px * slide_w))
            height = max(1, int(bh / h_px * slide_h))

            box = slide.shapes.add_textbox(left, top, width, height)
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

            p = tf.paragraphs[0]
            p.text = text
            p.font.size = Pt(_font_pt(reg))
            r, g, b = _rgb(reg)
            p.font.color.rgb = RGBColor(r, g, b)
            p.font.name = "Microsoft YaHei"

            ta = (reg.get("text_align") or "left").lower()
            if ta == "center":
                p.alignment = PP_ALIGN.CENTER
            elif ta == "right":
                p.alignment = PP_ALIGN.RIGHT
            elif ta == "justify":
                p.alignment = PP_ALIGN.JUSTIFY
            else:
                p.alignment = PP_ALIGN.LEFT

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
